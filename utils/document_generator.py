import os
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fpdf import FPDF

BASE_DIR = Path(__file__).resolve().parent.parent
ASSETS_DIR = BASE_DIR / "assets"
LOGO_PATH = ASSETS_DIR / "logo.png"


def sanitize_text(text: str) -> str:
    replacements = {
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "–": "-",
        "—": "-",
        "…": "...",
        "•": "-",
        "©": "(c)",
        "®": "(R)",
        "™": "TM",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode("latin-1", "replace").decode("latin-1")


def _set_cell_shading(cell, fill="E8EEF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def _extract_heading(line: str):
    clean = line.strip()
    if not clean:
        return None
    if re.match(r"^(SECTION\s+)?\d+[.)]\s+", clean, re.I):
        return re.sub(r"^(SECTION\s+)?\d+[.)]\s+", "", clean, flags=re.I)
    if clean.isupper() and len(clean) <= 120:
        return clean.title()
    return None


def _extract_terms(terms_text: str):
    return [item.strip() for item in re.split(r";|\n", terms_text) if item.strip()]


def format_docx(text: str, doc_type: str, terms_text: str = "") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = __import__("docx").shared.Inches(0.7)
    section.bottom_margin = __import__("docx").shared.Inches(0.7)
    section.left_margin = __import__("docx").shared.Inches(0.8)
    section.right_margin = __import__("docx").shared.Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = __import__("docx").shared.Pt(11)
    styles["Title"].font.name = "Times New Roman"
    styles["Title"].font.size = __import__("docx").shared.Pt(18)

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=__import__("docx").shared.Inches(1.2))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = __import__("docx").shared.Pt(18)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("AI-assisted draft - review before use").italic = True

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = _extract_heading(line)
        if heading:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = __import__("docx").shared.Pt(8)
            p.paragraph_format.space_after = __import__("docx").shared.Pt(4)
            r = p.add_run(heading)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = __import__("docx").shared.Pt(12)
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = __import__("docx").shared.Pt(5)
            p.paragraph_format.line_spacing = 1.15

    terms = _extract_terms(terms_text)
    if terms:
        doc.add_paragraph()
        heading = doc.add_paragraph()
        r = heading.add_run("USER-PROVIDED TERMS")
        r.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "No."
        hdr[1].text = "Term"
        for cell in hdr:
            _set_cell_shading(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for index, term in enumerate(terms, start=1):
            cells = table.add_row().cells
            cells[0].text = str(index)
            cells[1].text = term

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-assisted draft | Review by a qualified lawyer before signing | ")
    _add_page_number(footer)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = sanitize_text(doc_type)
        self.set_auto_page_break(auto=True, margin=18)

    def header(self):
        if LOGO_PATH.exists():
            try:
                self.image(str(LOGO_PATH), x=92, y=8, w=26)
                self.set_y(36)
            except Exception:
                self.set_y(10)
        else:
            self.set_y(10)
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 6, "LEGALEASE", align="C")
        self.ln(8)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 5, "LegalEase - AI-assisted draft | Review before signing | Page " + str(self.page_no()), align="C")


def format_pdf(text: str, doc_type: str, terms_text: str = "") -> bytes:
    pdf = LegalEasePDF(doc_type)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)

    pdf.set_font("Helvetica", "B", 16)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 9, sanitize_text(doc_type).upper(), align="C")
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(0, 6, "AI-assisted draft - review by a qualified lawyer before signing or use.", align="C")
    pdf.ln(5)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            pdf.ln(3)
            continue

        heading = _extract_heading(line)
        if heading:
            pdf.set_font("Helvetica", "B", 11)
            pdf.set_text_color(30, 40, 55)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 7, sanitize_text(heading))
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5.5, line)
            pdf.ln(1)

    terms = _extract_terms(terms_text)
    if terms:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(0, 7, "USER-PROVIDED TERMS")
        pdf.set_font("Helvetica", "", 9)
        for index, term in enumerate(terms, start=1):
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(0, 5, f"{index}. {sanitize_text(term)}")

    return bytes(pdf.output())


def format_txt(text: str) -> bytes:
    return text.encode("utf-8")
